from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import secrets
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from io import BytesIO, StringIO
from itertools import chain
from pathlib import Path
from typing import Literal
from uuid import uuid4

from dotenv import load_dotenv
from fastapi import FastAPI, File, Header, Query, Request, Response, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field
from sqlalchemy import create_engine, text
from sqlalchemy.engine import Connection

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")

DATABASE_URL = os.environ.get(
    "OTTORESTART_DATABASE_URL",
    "postgresql://restart_user:restart_secure_pass@127.0.0.1:5435/otto_restart_db",
)
MIGRATION_PATH = ROOT / "migrations" / "0001_postgres_initial.sql"
ACCOUNT = Literal["JV", "XL"]
PERIOD = Literal["3M", "12M"]
EAN_PATTERN = re.compile(r"^\d{8,14}$")
MAX_EANS_PER_REQUEST = 10_000

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,
    pool_size=5,
    max_overflow=10,
)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def format_iso(val: object) -> str:
    if hasattr(val, "isoformat"):
        return val.isoformat()
    return str(val)


@contextmanager
def connect():
    with engine.connect() as connection:
        yield connection


def migrate() -> None:
    migration_sql = MIGRATION_PATH.read_text(encoding="utf-8")
    with engine.connect() as conn:
        with conn.begin():
            conn.execute(text(migration_sql))
            conn.execute(
                text(
                    "INSERT INTO schema_migrations(version, applied_at) VALUES (1, CURRENT_TIMESTAMP) ON CONFLICT (version) DO NOTHING"
                )
            )


@asynccontextmanager
async def lifespan(_: FastAPI):
    migrate()
    yield


app = FastAPI(title="OttoRestart Local API", version="0.1.0", lifespan=lifespan)
allowed_origins = [
    origin.strip()
    for origin in os.environ.get(
        "OTTORESTART_ALLOWED_ORIGINS",
        "http://localhost:3001,http://127.0.0.1:3001,http://localhost:3002,http://127.0.0.1:3002",
    ).split(",")
    if origin.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type", "X-CSRF-Token", "Idempotency-Key"],
)


@app.get("/api/restart/health")
def health():
    with connect() as conn:
        conn.execute(text("SELECT 1"))
    return {"status": "ok", "service": "ottorestart-local-api", "database": "postgresql"}


class BatchInput(BaseModel):
    account: ACCOUNT
    period: PERIOD
    name: str = Field(min_length=1, max_length=160)
    sourceEans: list[str] = Field(min_length=1, max_length=MAX_EANS_PER_REQUEST)


class EanImportInput(BaseModel):
    account: ACCOUNT
    eans: list[str] = Field(min_length=1, max_length=MAX_EANS_PER_REQUEST)


def normalize_ean(raw: str) -> str:
    return raw.strip().replace("\ufeff", "")


def valid_ean(raw: str) -> bool:
    if not EAN_PATTERN.fullmatch(raw):
        return False
    digits = [int(char) for char in raw]
    check_digit = digits.pop()
    weighted_sum = sum(digit * (3 if index % 2 == 0 else 1) for index, digit in enumerate(reversed(digits)))
    return (10 - weighted_sum % 10) % 10 == check_digit


def inspect_eans(raw_eans: list[str]) -> dict:
    normalized = [normalize_ean(value) for value in raw_eans]
    seen: set[str] = set()
    duplicates: list[str] = []
    invalid: list[str] = []
    unique: list[str] = []
    for value in normalized:
        if not valid_ean(value):
            invalid.append(value)
            continue
        if value in seen:
            duplicates.append(value)
            continue
        seen.add(value)
        unique.append(value)
    return {
        "inputCount": len(raw_eans),
        "uniqueCount": len(unique),
        "duplicateCount": len(duplicates),
        "duplicates": duplicates,
        "invalidEans": invalid,
        "uniqueEans": unique,
    }


def parse_delimited_text(content: str) -> list[str]:
    nonempty_lines = [line for line in content.splitlines() if line.strip()]
    if not nonempty_lines:
        return []
    sample = "\n".join(nonempty_lines[:10])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(StringIO("\n".join(nonempty_lines)), dialect))
    if not rows:
        return []
    header = [cell.strip().casefold() for cell in rows[0]]
    ean_headers = {"ean", "ean13", "gtin", "barcode", "штрихкод"}
    column = next((index for index, value in enumerate(header) if value in ean_headers), 0)
    start = 1 if any(value in ean_headers for value in header) else 0
    return [row[column].strip() for row in rows[start:] if len(row) > column and row[column].strip()]


def parse_xlsx(content: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ApiProblem(503, "INTERNAL_ERROR", "XLSX support is not installed in this OttoRestart backend.") from error
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows()
        header_row = next(rows, None)
    except Exception as error:
        raise ApiProblem(422, "VALIDATION_ERROR", "Could not read this XLSX file.") from error
    if header_row is None:
        workbook.close()
        return []
    header = [str(cell.value or "").strip().casefold() for cell in header_row]
    ean_headers = {"ean", "ean13", "gtin", "barcode", "штрихкод"}
    column = next((index for index, value in enumerate(header) if value in ean_headers), 0)
    has_header = any(value in ean_headers for value in header)
    output: list[str] = []
    data_rows = rows if has_header else chain([header_row], rows)
    for row_number, row in enumerate(data_rows, start=1):
        if row_number > MAX_EANS_PER_REQUEST:
            workbook.close()
            raise ApiProblem(422, "VALIDATION_ERROR", f"A file may contain at most {MAX_EANS_PER_REQUEST} EAN rows.")
        if len(row) <= column:
            continue
        cell = row[column]
        value = cell.value
        if value is None:
            continue
        if isinstance(value, bool):
            continue
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        text_val = str(value).strip()
        if isinstance(value, int) and re.fullmatch(r"0+", cell.number_format or ""):
            text_val = text_val.zfill(len(cell.number_format))
        output.append(text_val)
    workbook.close()
    return output


async def read_ean_upload(file: UploadFile) -> list[str]:
    filename = (file.filename or "").lower()
    content = await file.read(20_000_001)
    if len(content) > 20_000_000:
        raise ApiProblem(413, "VALIDATION_ERROR", "The import file exceeds the 20 MB limit.")
    if filename.endswith(".xlsx"):
        values = parse_xlsx(content)
    elif filename.endswith(".csv") or filename.endswith(".txt"):
        values = parse_delimited_text(content.decode("utf-8-sig", errors="replace"))
    else:
        raise ApiProblem(422, "VALIDATION_ERROR", "Use an .xlsx, .csv, or .txt file.")
    if len(values) > MAX_EANS_PER_REQUEST:
        raise ApiProblem(422, "VALIDATION_ERROR", f"A file may contain at most {MAX_EANS_PER_REQUEST} EAN rows.")
    return values


def page_info(page: int, page_size: int, total_items: int) -> dict:
    return {
        "page": page,
        "pageSize": page_size,
        "totalItems": total_items,
        "totalPages": (total_items + page_size - 1) // page_size if total_items > 0 else 0,
    }


def batch_dict(row: dict) -> dict:
    return {
        "id": row["id"],
        "account": row["account"],
        "period": row["period"],
        "name": row["name"],
        "status": row["status"],
        "itemCount": row["item_count"],
        "createdAt": format_iso(row["created_at"]),
    }


def add_event(
    conn: Connection,
    account: str,
    event_type: str,
    entity_type: str,
    entity_id: str,
    summary: str,
    metadata: dict,
):
    conn.execute(
        text(
            """
            INSERT INTO restart_events(id, account, event_type, entity_type, entity_id, summary, metadata_json, created_at)
            VALUES (:id, :account, :event_type, :entity_type, :entity_id, :summary, :metadata_json, CURRENT_TIMESTAMP)
            """
        ),
        {
            "id": str(uuid4()),
            "account": account,
            "event_type": event_type,
            "entity_type": entity_type,
            "entity_id": entity_id,
            "summary": summary,
            "metadata_json": json.dumps(metadata, ensure_ascii=False),
        },
    )


def csrf_cookie(request: Request) -> str | None:
    return request.cookies.get("ottorestart_csrf")


def require_local_write(request: Request, csrf_header: str | None) -> None:
    origin = request.headers.get("origin")
    if origin not in allowed_origins:
        raise ApiProblem(403, "FORBIDDEN", "Write requests are limited to configured local origins.")
    cookie = csrf_cookie(request)
    if not cookie or not csrf_header or not secrets.compare_digest(cookie, csrf_header):
        raise ApiProblem(403, "FORBIDDEN", "Missing or invalid OttoRestart CSRF token.")


class ApiProblem(Exception):
    def __init__(self, status: int, code: str, message: str, details: dict | None = None):
        self.status = status
        self.code = code
        self.message = message
        self.details = details


@app.exception_handler(ApiProblem)
async def handle_api_problem(_: Request, error: ApiProblem):
    return JSONResponse(
        status_code=error.status,
        content={
            "code": error.code,
            "message": error.message,
            "requestId": str(uuid4()),
            **({"details": error.details} if error.details else {}),
        },
    )


@app.exception_handler(RequestValidationError)
async def handle_validation_error(_: Request, error: RequestValidationError):
    return JSONResponse(
        status_code=422,
        content={
            "code": "VALIDATION_ERROR",
            "message": "The request did not match the OttoRestart API contract.",
            "requestId": str(uuid4()),
            "details": {"errors": jsonable_encoder(error.errors())},
        },
    )


@app.get("/api/restart/csrf")
def get_csrf(response: Response):
    token = secrets.token_urlsafe(32)
    response.set_cookie("ottorestart_csrf", token, httponly=False, samesite="strict", secure=False, path="/")
    return {"token": token}


@app.get("/api/restart/overview")
def get_overview(account: ACCOUNT):
    with connect() as conn:
        batches = conn.execute(
            text(
                "SELECT COALESCE(SUM(item_count), 0) FROM restart_batches WHERE account = :account AND status != 'CANCELLED'"
            ),
            {"account": account},
        ).scalar() or 0
        eans = conn.execute(
            text("SELECT COUNT(*) FROM restart_ean_registry WHERE account = :account AND status = 'AVAILABLE'"),
            {"account": account},
        ).scalar() or 0
    return {
        "account": account,
        "counts": {
            "candidates": batches,
            "awaitingReview": 0,
            "needsAttention": 0,
            "queued": 0,
            "onlineConfirmed": 0,
            "availableEans": eans,
        },
        "generatedAt": utc_now(),
    }


@app.get("/api/restart/batches")
def list_batches(
    account: ACCOUNT,
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=200),
    status: str | None = None,
    query: str | None = Query(None, max_length=200),
):
    where = ["account = :account"]
    params: dict[str, object] = {"account": account}
    if status:
        where.append("status = :status")
        params["status"] = status
    if query:
        where.append("name ILIKE :query")
        params["query"] = f"%{query}%"
    where_sql = " AND ".join(where)

    with connect() as conn:
        total = conn.execute(
            text(f"SELECT COUNT(*) FROM restart_batches WHERE {where_sql}"),
            params,
        ).scalar() or 0

        params["limit"] = pageSize
        params["offset"] = (page - 1) * pageSize
        rows = conn.execute(
            text(f"SELECT * FROM restart_batches WHERE {where_sql} ORDER BY created_at DESC LIMIT :limit OFFSET :offset"),
            params,
        ).mappings().all()

    return {"items": [batch_dict(dict(row)) for row in rows], "page": page_info(page, pageSize, total)}


@app.post("/api/restart/batches/preview")
def preview_batch(
    payload: BatchInput,
    request: Request,
    x_csrf_token: str | None = Header(None, alias="X-CSRF-Token"),
):
    require_local_write(request, x_csrf_token)
    if not payload.name.strip():
        raise ApiProblem(422, "VALIDATION_ERROR", "Batch name cannot be blank.")
    inspection = inspect_eans(payload.sourceEans)
    return {
        "account": payload.account,
        "period": payload.period,
        "inputCount": inspection["inputCount"],
        "uniqueCount": inspection["uniqueCount"],
        "duplicateCount": inspection["duplicateCount"],
        "duplicates": inspection["duplicates"],
        "invalidEans": inspection["invalidEans"],
        "sourceEans": inspection["uniqueEans"],
        "warnings": ["Проверен только формат EAN и контрольная цифра; внешние источники не вызывались."],
        "dataSourcesChecked": ["LOCAL_FORMAT_ONLY"],
    }


@app.post("/api/restart/imports/preview")
async def preview_import_file(
    request: Request,
    account: ACCOUNT = Query(...),
    period: PERIOD | None = Query(None),
    file: UploadFile = File(...),
    x_csrf_token: str | None = Header(None, alias="X-CSRF-Token"),
):
    require_local_write(request, x_csrf_token)
    values = await read_ean_upload(file)
    inspection = inspect_eans(values)
    return {
        "account": account,
        "period": period,
        "inputCount": inspection["inputCount"],
        "uniqueCount": inspection["uniqueCount"],
        "duplicateCount": inspection["duplicateCount"],
        "duplicates": inspection["duplicates"],
        "invalidEans": inspection["invalidEans"],
        "sourceEans": inspection["uniqueEans"],
        "warnings": ["Проверен только файл и контрольная цифра EAN; внешние источники не вызывались."],
        "dataSourcesChecked": ["LOCAL_FORMAT_ONLY"],
    }


@app.post("/api/restart/batches", status_code=201)
def create_batch(
    payload: BatchInput,
    request: Request,
    idempotency_key: str = Header(..., alias="Idempotency-Key", min_length=8, max_length=128),
    x_csrf_token: str | None = Header(None, alias="X-CSRF-Token"),
):
    require_local_write(request, x_csrf_token)
    if not payload.name.strip():
        raise ApiProblem(422, "VALIDATION_ERROR", "Batch name cannot be blank.")
    inspection = inspect_eans(payload.sourceEans)
    if inspection["invalidEans"] or inspection["duplicates"]:
        raise ApiProblem(
            422,
            "VALIDATION_ERROR",
            "Batch contains invalid or duplicate EANs; run preview and correct the input.",
            {"invalidEans": inspection["invalidEans"], "duplicates": inspection["duplicates"]},
        )
    canonical = payload.model_dump(mode="json")
    body_hash = hashlib.sha256(json.dumps(canonical, sort_keys=True).encode("utf-8")).hexdigest()
    created_at = utc_now()
    batch_id = str(uuid4())

    with connect() as conn:
        with conn.begin():
            previous = conn.execute(
                text("SELECT request_hash, response_body FROM restart_idempotency_keys WHERE key = :key FOR UPDATE"),
                {"key": idempotency_key},
            ).mappings().first()
            if previous:
                if previous["request_hash"] != body_hash:
                    raise ApiProblem(409, "CONFLICT", "Idempotency key was already used for a different request.")
                return json.loads(previous["response_body"])

            conn.execute(
                text(
                    """
                    INSERT INTO restart_batches(id, account, period, name, status, item_count, created_at)
                    VALUES (:id, :account, :period, :name, :status, :item_count, CURRENT_TIMESTAMP)
                    """
                ),
                {
                    "id": batch_id,
                    "account": payload.account,
                    "period": payload.period,
                    "name": payload.name.strip(),
                    "status": "IMPORTED",
                    "item_count": len(inspection["uniqueEans"]),
                },
            )

            conn.execute(
                text(
                    """
                    INSERT INTO restart_items(batch_id, account, source_ean, row_number)
                    VALUES (:batch_id, :account, :source_ean, :row_number)
                    """
                ),
                [
                    {
                        "batch_id": batch_id,
                        "account": payload.account,
                        "source_ean": ean,
                        "row_number": index,
                    }
                    for index, ean in enumerate(inspection["uniqueEans"], start=1)
                ],
            )

            result = {
                "id": batch_id,
                "account": payload.account,
                "period": payload.period,
                "name": payload.name.strip(),
                "status": "IMPORTED",
                "itemCount": len(inspection["uniqueEans"]),
                "createdAt": created_at,
            }

            conn.execute(
                text(
                    """
                    INSERT INTO restart_idempotency_keys(key, request_hash, response_body, created_at)
                    VALUES (:key, :request_hash, :response_body, CURRENT_TIMESTAMP)
                    """
                ),
                {
                    "key": idempotency_key,
                    "request_hash": body_hash,
                    "response_body": json.dumps(result),
                },
            )

            add_event(
                conn,
                payload.account,
                "BATCH_IMPORTED",
                "BATCH",
                batch_id,
                f"Импортирована партия «{payload.name.strip()}»",
                {"period": payload.period, "itemCount": result["itemCount"]},
            )

    return result


@app.get("/api/restart/batches/{batch_id}")
def get_batch(batch_id: str, account: ACCOUNT):
    with connect() as conn:
        row = conn.execute(
            text("SELECT * FROM restart_batches WHERE id = :id AND account = :account"),
            {"id": batch_id, "account": account},
        ).mappings().first()
    if not row:
        raise ApiProblem(404, "NOT_FOUND", "Batch was not found.")
    return batch_dict(dict(row))


@app.get("/api/restart/batches/{batch_id}/items")
def list_batch_items(
    batch_id: str,
    account: ACCOUNT,
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=200),
    query: str | None = Query(None, max_length=100),
):
    with connect() as conn:
        exists = conn.execute(
            text("SELECT 1 FROM restart_batches WHERE id = :id AND account = :account"),
            {"id": batch_id, "account": account},
        ).first()
        if not exists:
            raise ApiProblem(404, "NOT_FOUND", "Batch was not found for this account.")

        where = "batch_id = :batch_id AND account = :account"
        params: dict[str, object] = {"batch_id": batch_id, "account": account}
        if query:
            where += " AND source_ean ILIKE :query"
            params["query"] = f"%{query}%"

        total = conn.execute(
            text(f"SELECT COUNT(*) FROM restart_items WHERE {where}"),
            params,
        ).scalar() or 0

        params["limit"] = pageSize
        params["offset"] = (page - 1) * pageSize
        rows = conn.execute(
            text(
                f"""
                SELECT id, batch_id, account, source_ean, row_number
                FROM restart_items
                WHERE {where}
                ORDER BY row_number
                LIMIT :limit OFFSET :offset
                """
            ),
            params,
        ).mappings().all()

    return {
        "items": [
            {
                "id": str(row["id"]),
                "batchId": row["batch_id"],
                "account": row["account"],
                "sourceEan": row["source_ean"],
                "rowNumber": row["row_number"],
            }
            for row in rows
        ],
        "page": page_info(page, pageSize, total),
    }


@app.get("/api/restart/ean-pool")
def list_ean_pool(
    account: ACCOUNT,
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=200),
    status: str | None = None,
):
    where = ["account = :account"]
    params: dict[str, object] = {"account": account}
    if status:
        where.append("status = :status")
        params["status"] = status
    where_sql = " AND ".join(where)

    with connect() as conn:
        total = conn.execute(
            text(f"SELECT COUNT(*) FROM restart_ean_registry WHERE {where_sql}"),
            params,
        ).scalar() or 0

        params["limit"] = pageSize
        params["offset"] = (page - 1) * pageSize
        rows = conn.execute(
            text(
                f"""
                SELECT * FROM restart_ean_registry
                WHERE {where_sql}
                ORDER BY imported_at DESC, ean
                LIMIT :limit OFFSET :offset
                """
            ),
            params,
        ).mappings().all()

    return {
        "items": [
            {
                "ean": row["ean"],
                "account": row["account"],
                "status": row["status"],
                "batchId": row["batch_id"],
                "importedAt": format_iso(row["imported_at"]),
            }
            for row in rows
        ],
        "page": page_info(page, pageSize, total),
    }


@app.post("/api/restart/ean-pool/import", status_code=201)
def import_eans(
    payload: EanImportInput,
    request: Request,
    x_csrf_token: str | None = Header(None, alias="X-CSRF-Token"),
):
    require_local_write(request, x_csrf_token)
    inspection = inspect_eans(payload.eans)
    if inspection["invalidEans"] or inspection["duplicates"]:
        raise ApiProblem(
            422,
            "VALIDATION_ERROR",
            "EAN import contains invalid or duplicate values.",
            {"invalidEans": inspection["invalidEans"], "duplicates": inspection["duplicates"]},
        )
    imported_at = utc_now()
    eans = inspection["uniqueEans"]

    with connect() as conn:
        with conn.begin():
            existing = []
            for start in range(0, len(eans), 500):
                chunk = eans[start : start + 500]
                rows = conn.execute(
                    text("SELECT ean FROM restart_ean_registry WHERE ean = ANY(:eans)"),
                    {"eans": chunk},
                ).mappings().all()
                existing.extend(r["ean"] for r in rows)

            if existing:
                raise ApiProblem(
                    409,
                    "CONFLICT",
                    "Some EANs already exist in the OttoRestart pool.",
                    {"existingEans": existing},
                )

            conn.execute(
                text(
                    """
                    INSERT INTO restart_ean_registry(ean, account, status, batch_id, imported_at)
                    VALUES (:ean, :account, 'AVAILABLE', NULL, CURRENT_TIMESTAMP)
                    """
                ),
                [{"ean": ean, "account": payload.account} for ean in eans],
            )

            add_event(
                conn,
                payload.account,
                "EAN_POOL_IMPORTED",
                "EAN_POOL",
                imported_at,
                f"Добавлено EAN в пул: {len(eans)}",
                {"count": len(eans)},
            )

    return {
        "importedCount": len(eans),
        "account": payload.account,
        "status": "AVAILABLE",
        "importedAt": imported_at,
    }


@app.get("/api/restart/events")
def list_events(
    account: ACCOUNT,
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=200),
):
    with connect() as conn:
        total = conn.execute(
            text("SELECT COUNT(*) FROM restart_events WHERE account = :account"),
            {"account": account},
        ).scalar() or 0

        rows = conn.execute(
            text(
                """
                SELECT * FROM restart_events
                WHERE account = :account
                ORDER BY created_at DESC
                LIMIT :limit OFFSET :offset
                """
            ),
            {"account": account, "limit": pageSize, "offset": (page - 1) * pageSize},
        ).mappings().all()

    return {
        "items": [
            {
                "id": str(row["id"]),
                "account": row["account"],
                "eventType": row["event_type"],
                "entityType": row["entity_type"],
                "entityId": row["entity_id"],
                "summary": row["summary"],
                "metadata": row["metadata_json"] if isinstance(row["metadata_json"], dict) else json.loads(row["metadata_json"]),
                "createdAt": format_iso(row["created_at"]),
            }
            for row in rows
        ],
        "page": page_info(page, pageSize, total),
    }


@app.get("/api/restart/queue")
def get_queue(
    account: ACCOUNT,
    page: int = Query(1, ge=1),
    pageSize: int = Query(50, ge=1, le=200),
):
    return {
        "account": account,
        "status": "IDLE",
        "items": [],
        "page": page_info(page, pageSize, 0),
    }
